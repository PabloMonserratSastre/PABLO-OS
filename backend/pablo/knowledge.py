import math
import re
from collections import Counter
from io import BytesIO
from zipfile import ZipFile

from sqlalchemy import select

from .db import Chunk, Item, Owner
from .providers import CompatibleProvider


def extract(name: str, data: bytes) -> str:
    ext = name.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        from .pdf_tables import extract_pdf
        text = extract_pdf(data)
    elif ext == "docx":
        with ZipFile(BytesIO(data)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 20_000_000:
                raise ValueError("Documento expandido demasiado grande.")
        from docx import Document

        doc = Document(BytesIO(data))
        text = "\n".join(
            [p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells]
        )
    elif ext in {"txt", "md", "py", "js", "ts", "tsx", "cs", "json", "csv", "sql"}:
        text = data.decode("utf-8")
    else:
        raise ValueError("Formato admitido: PDF, DOCX, TXT, Markdown o código UTF-8.")
    if not text.strip():
        raise ValueError("No se ha encontrado texto. Los PDF escaneados necesitan OCR externo.")
    if len(text) > 500_000:
        raise ValueError("Máximo 500.000 caracteres extraídos.")
    return text


def extract_isolated(name: str, data: bytes) -> str:
    extension = name.rsplit(".", 1)[-1].lower()
    if extension not in {"pdf", "docx"}:
        return extract(name, data)
    import json
    import os
    import subprocess
    import sys

    options = {"creationflags": subprocess.CREATE_NO_WINDOW} if os.name == "nt" else {}
    try:
        result = subprocess.run([sys.executable, "-m", "pablo.extractor", extension], input=data,
                                capture_output=True, timeout=20, check=True, **options)
        parsed = json.loads(result.stdout)
    except (subprocess.SubprocessError, ValueError):
        raise ValueError("El documento excedió el tiempo de extracción o no pudo procesarse.")
    if parsed.get("error"):
        raise ValueError(parsed["error"])
    return parsed["text"]


def remote_allowed(db):
    owner = db.get(Owner, 1)
    return bool(owner and not owner.settings.get("demo", False))


def pieces_for(text):
    from .pdf_tables import MARKER
    if text.startswith(MARKER):
        return [block for block in text.split('\n\n') if block.strip()]
    return [text[i : i + 1200] for i in range(0, len(text), 1000)]


def restore_text(chunks):
    from .pdf_tables import MARKER
    if chunks and chunks[0].text.startswith(MARKER):
        return '\n\n'.join(row.text for row in chunks)
    return ''.join(row.text[:1000] for row in chunks[:-1]) + (chunks[-1].text if chunks else '')


def prepare_vectors(text: str):
    pieces = pieces_for(text)
    vectors = []
    from .db import DB

    with DB() as db:
        if not remote_allowed(db):
            return vectors
    provider = CompatibleProvider()
    if not provider.config["embed_model"] or not provider.key:
        return vectors
    for start in range(0, len(pieces), 24):
        batch = provider.embed(pieces[start : start + 24])
        vectors.extend(batch)
    return vectors


def index(db, document_id: str, text: str, vectors=None):
    pieces = pieces_for(text)
    vectors = vectors or []
    for i, piece in enumerate(pieces):
        db.add(
            Chunk(
                document_id=document_id,
                position=i,
                text=piece,
                embedding=vectors[i] if len(vectors) == len(pieces) else None,
            )
        )
    return len(pieces)


def tokens(text: str) -> Counter:
    from .pdf_tables import normalized
    text = normalized(text)
    text = re.sub(r'\bs1\b', 'primer cuatrimestre', text)
    text = re.sub(r'\bs2\b', 'segundo cuatrimestre', text)
    text = re.sub(r'\bcuatri\b', 'cuatrimestre', text)
    stop = {
        "de",
        "el",
        "la",
        "los",
        "las",
        "un",
        "una",
        "que",
        "qué",
        "en",
        "y",
        "por",
        "para",
        "del",
        "mis",
        "me",
        "documento",
        "documentos",
        "dime",
    }
    return Counter(w for w in re.findall(r"\w+", text) if len(w) > 2 and w not in stop)


def search(db, query: str, project_id: str | None = None) -> list[dict]:
    statement = select(Chunk, Item).join(Item, Item.id == Chunk.document_id)
    if project_id:
        statement = statement.where(Item.data["project_id"].as_string() == project_id)
    candidates = db.execute(statement.limit(4000)).all()
    q = tokens(query)
    embedded = (
        CompatibleProvider().embed([query])
        if remote_allowed(db) and any(chunk.embedding for chunk, _ in candidates)
        else []
    )
    ranked = []
    for chunk, doc in candidates:
        words = tokens(chunk.text)
        score = sum(min(count, words[word]) for word, count in q.items()) / max(
            1, math.sqrt(sum(words.values()))
        )
        method = "lexical"
        if embedded and chunk.embedding and len(embedded[0]) == len(chunk.embedding):
            a, b = embedded[0], chunk.embedding
            score = sum(x * y for x, y in zip(a, b)) / max(
                1e-12, math.sqrt(sum(x * x for x in a) * sum(x * x for x in b))
            )
            method = "semantic"
        if score > 0:
            ranked.append(
                {
                    "document_id": doc.id,
                    "source": doc.title,
                    "fragment": chunk.position + 1,
                    "text": chunk.text,
                    "score": round(score, 4),
                    "method": method,
                }
            )
    return sorted(ranked, key=lambda row: row["score"], reverse=True)[:6]
